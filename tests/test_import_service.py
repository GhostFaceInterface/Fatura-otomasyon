from pathlib import Path

import pandas as pd
import pytest

from invoice_automation.app.db.repository import InvoiceRecordRepository
from invoice_automation.app.services.import_service import ImportService, parse_currency
from invoice_automation.app.utils.exceptions import ImportValidationError


def test_import_service_imports_valid_csv_rows(tmp_path: Path) -> None:
    csv_path = tmp_path / "valid.csv"
    csv_path.write_text(
        "ad,soyad,tc_kimlik_no,tutar_usd,aciklama\n"
        "Ali,Yilmaz,12345678901,1500,Yurt disi konaklama\n",
        encoding="utf-8",
    )
    repository = InvoiceRecordRepository(tmp_path / "test.sqlite3")

    result = ImportService(repository).import_file(csv_path)

    assert result.imported_count == 1
    assert result.failed_count == 0
    assert repository.count() == 1
    assert repository.list_all()[0].kaynak_dosya == "valid.csv"
    assert result.batch is not None
    assert repository.list_all()[0].batch_id == result.batch.id


@pytest.mark.parametrize(
    ("raw_value", "expected"),
    [
        ("$1.450", 1450.0),
        ("$850", 850.0),
        ("1.350", 1350.0),
        ("1350", 1350.0),
        ("$1,450", 1450.0),
        ("$1450.75", 1450.75),
    ],
)
def test_parse_currency_normalizes_import_amounts(raw_value: str, expected: float) -> None:
    assert parse_currency(raw_value) == expected


@pytest.mark.parametrize("raw_value", ["", "   ", None, "abc"])
def test_parse_currency_rejects_empty_or_unparseable_values(raw_value) -> None:
    with pytest.raises(ValueError):
        parse_currency(raw_value)


def test_import_service_normalizes_currency_amounts_before_persisting(tmp_path: Path) -> None:
    csv_path = tmp_path / "amounts.csv"
    csv_path.write_text(
        "ad,soyad,tc_kimlik_no,tutar_usd,aciklama\n"
        'Ali,Yilmaz,12345678901,"$1.450",Gecerli\n'
        'Ayse,Demir,10987654321,"$1,450",Gecerli\n',
        encoding="utf-8",
    )
    repository = InvoiceRecordRepository(tmp_path / "test.sqlite3")

    result = ImportService(repository).import_file(csv_path)
    records = repository.list_all(batch_id=result.batch.id)

    assert result.imported_count == 2
    assert result.failed_count == 0
    assert sorted(record.tutar_usd for record in records) == [1450.0, 1450.0]


def test_import_service_rejects_missing_columns(tmp_path: Path) -> None:
    csv_path = tmp_path / "missing.csv"
    csv_path.write_text("ad,soyad\nAli,Yilmaz\n", encoding="utf-8")
    repository = InvoiceRecordRepository(tmp_path / "test.sqlite3")

    with pytest.raises(ImportValidationError):
        ImportService(repository).import_file(csv_path)


def test_import_service_skips_invalid_rows_without_crashing(tmp_path: Path) -> None:
    csv_path = tmp_path / "mixed.csv"
    csv_path.write_text(
        "ad,soyad,tc_kimlik_no,tutar_usd,aciklama\n"
        "Ali,Yilmaz,12345678901,1500,Gecerli\n"
        "Ayse,Demir,invalid,2200,Hatali tckn\n"
        "Mehmet,Kaya,10987654321,-10,Hatali tutar\n",
        encoding="utf-8",
    )
    repository = InvoiceRecordRepository(tmp_path / "test.sqlite3")

    result = ImportService(repository).import_file(csv_path)

    assert result.imported_count == 1
    assert result.failed_count == 2
    assert repository.count() == 1
    assert [error.row_number for error in result.row_errors] == [3, 4]


def test_import_service_imports_selected_excel_sheet_with_column_mapping(tmp_path: Path) -> None:
    excel_path = tmp_path / "payments.xlsx"
    with pd.ExcelWriter(excel_path) as writer:
        pd.DataFrame({"ignored": ["x"]}).to_excel(writer, sheet_name="Bos", index=False)
        pd.DataFrame(
            {
                "Adi": ["Ali"],
                "Soyadi": ["Yilmaz"],
                "Kimlik": ["12345678901"],
                "Odenen": ["1500"],
            }
        ).to_excel(writer, sheet_name="Nisan", index=False)
    repository = InvoiceRecordRepository(tmp_path / "test.sqlite3")

    inspection = ImportService(repository).inspect_file(excel_path, sheet_name="Nisan")
    result = ImportService(repository).import_file(
        excel_path,
        batch_name="2026-04 Nisan",
        sheet_name="Nisan",
        column_mapping={
            "ad": "Adi",
            "soyad": "Soyadi",
            "tc_kimlik_no": "Kimlik",
            "tutar_usd": "Odenen",
            "aciklama": "",
        },
    )

    records = repository.list_all(batch_id=result.batch.id)
    assert inspection.sheet_names == ["Bos", "Nisan"]
    assert inspection.columns == ["Adi", "Soyadi", "Kimlik", "Odenen"]
    assert result.batch.name == "2026-04 Nisan"
    assert result.batch.sheet_name == "Nisan"
    assert len(records) == 1
    assert records[0].ad == "Ali"


@pytest.mark.parametrize(
    "raw_value",
    ["NaN", "nan", "sNaN", "Infinity", "-Infinity", "$inf", float("inf"), float("-inf")],
)
def test_parse_currency_rejects_non_finite_values(raw_value) -> None:
    with pytest.raises(ValueError):
        parse_currency(raw_value)


def test_import_service_reports_non_finite_amount_as_row_error(tmp_path: Path) -> None:
    csv_path = tmp_path / "nan.csv"
    csv_path.write_text(
        "ad,soyad,tc_kimlik_no,tutar_usd\n"
        "Ali,Yilmaz,12345678901,NaN\n"
        "Ayse,Demir,10987654321,Infinity\n"
        "Veli,Kaya,11111111111,150\n",
        encoding="utf-8",
    )
    repository = InvoiceRecordRepository(tmp_path / "test.sqlite3")

    result = ImportService(repository).import_file(csv_path)

    assert result.imported_count == 1
    assert result.failed_count == 2
    assert repository.count() == 1


def test_import_service_rolls_back_whole_import_on_persist_failure(tmp_path: Path, monkeypatch) -> None:
    csv_path = tmp_path / "valid.csv"
    csv_path.write_text(
        "ad,soyad,tc_kimlik_no,tutar_usd\n"
        "Ali,Yilmaz,12345678901,100\n"
        "Ayse,Demir,10987654321,200\n",
        encoding="utf-8",
    )
    repository = InvoiceRecordRepository(tmp_path / "test.sqlite3")
    from invoice_automation.app.db import repository as repository_module

    original_insert = repository_module._insert_record
    calls = {"count": 0}

    def failing_insert(connection, record, now):
        calls["count"] += 1
        if calls["count"] == 2:
            raise RuntimeError("disk full")
        return original_insert(connection, record, now)

    monkeypatch.setattr(repository_module, "_insert_record", failing_insert)

    with pytest.raises(RuntimeError):
        ImportService(repository).import_file(csv_path)

    assert repository.count() == 0
    assert repository.list_import_batches() == []


def _write_titled_payment_sheet(excel_path: Path) -> None:
    """Write a sheet shaped like the real payment lists: title rows, header on row 4."""

    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Odemeler"
    sheet["A1"] = "6 EKIM 2026 ODEMELER"
    sheet.merge_cells("A1:F1")
    sheet["A2"] = "GIDIS TARIHI :"
    sheet["D2"] = "23.03.2026"
    sheet["G2"] = "ACIKLAMA"
    sheet.merge_cells("G2:G4")
    sheet["A3"] = "DONUS TARIHI :"
    sheet["D3"] = "11.04.2026"
    for column, name in enumerate(["NO", "ADI", "SOYADI", "TC", "UCRET", "ODENEN"], start=1):
        sheet.cell(4, column, name)
    sheet.cell(5, 1, 1)
    sheet.cell(5, 2, "Ali")
    sheet.cell(5, 3, "Yilmaz")
    sheet.cell(5, 4, 12345678901)
    sheet.cell(5, 5, 2100)
    sheet.cell(5, 6, 2100)
    sheet.cell(5, 7, "Not")
    sheet.cell(7, 5, 2100)  # totals row under an empty line
    workbook.save(excel_path)


def test_inspect_file_finds_header_below_title_rows(tmp_path: Path) -> None:
    excel_path = tmp_path / "odemeler.xlsx"
    _write_titled_payment_sheet(excel_path)

    inspection = ImportService(InvoiceRecordRepository(tmp_path / "test.sqlite3")).inspect_file(excel_path)

    assert inspection.columns == ["NO", "ADI", "SOYADI", "TC", "UCRET", "ODENEN", "ACIKLAMA"]


def test_import_file_reads_rows_below_title_rows(tmp_path: Path) -> None:
    excel_path = tmp_path / "odemeler.xlsx"
    _write_titled_payment_sheet(excel_path)
    repository = InvoiceRecordRepository(tmp_path / "test.sqlite3")

    result = ImportService(repository).import_file(
        excel_path,
        sheet_name="Odemeler",
        column_mapping={
            "ad": "ADI",
            "soyad": "SOYADI",
            "tc_kimlik_no": "TC",
            "tutar_usd": "UCRET",
            "aciklama": "ACIKLAMA",
        },
    )

    records = repository.list_all(batch_id=result.batch.id)
    assert [(record.ad, record.tc_kimlik_no, record.kaynak_satir_no) for record in records] == [
        ("Ali", "12345678901", 5)
    ]
    assert records[0].aciklama == "Not"
    assert [error.row_number for error in result.row_errors] == [7]
