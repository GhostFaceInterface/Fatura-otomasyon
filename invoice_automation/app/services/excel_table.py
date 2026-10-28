"""Read an Excel sheet whose column names are not on the first row.

Real payment lists often start with a merged title row and a few date rows,
and some column names sit in cells merged down into the header row. This
module finds the header row and returns the table below it.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any, Callable

import pandas as pd

# How many rows from the top are searched for the header row.
HEADER_SEARCH_ROWS = 30

# Column names (after normalize_column_name) that mark a row as the header.
KNOWN_COLUMNS = frozenset({"ad", "soyad", "tc_kimlik_no", "tutar_usd", "aciklama"})


def read_excel_table(
    file_path: Path,
    sheet_name: str | None,
    normalize: Callable[[str], str],
) -> pd.DataFrame:
    """Return the sheet's table with its real column names.

    The index keeps sheet positions, so ``index + 2`` is the Excel row number
    of a data row, as it is for a sheet whose header is on the first row.
    Columns without a name and fully empty rows are left out.
    """

    raw = pd.read_excel(
        file_path,
        sheet_name=sheet_name if sheet_name is not None else 0,
        header=None,
        dtype=str,
        keep_default_na=False,
    )
    if raw.empty:
        return pd.DataFrame()

    header_index = find_header_row(raw, normalize)
    names = _header_names(raw, header_index, _merged_header_names(file_path, sheet_name, header_index))

    data = raw.iloc[header_index + 1 :]
    keep = [position for position, name in enumerate(names) if name]
    table = data.iloc[:, keep].copy()
    table.columns = _unique([names[position] for position in keep])
    table.index = [int(index) - 1 for index in data.index]
    has_value = table.apply(lambda row: any(_text(value) for value in row), axis=1)
    return table[has_value] if len(table) else table


def find_header_row(raw: pd.DataFrame, normalize: Callable[[str], str]) -> int:
    """Return the position of the row that holds the column names.

    The row with the most known column names wins. Without at least two of
    them, the first row with two or more filled cells is used, which skips a
    lone title row.
    """

    rows = range(min(len(raw), HEADER_SEARCH_ROWS))
    best_index, best_score = 0, 0
    for index in rows:
        found = {normalize(_text(value)) for value in raw.iloc[index] if _text(value)}
        score = len(found & KNOWN_COLUMNS)
        if score > best_score:
            best_index, best_score = index, score
    if best_score >= 2:
        return best_index

    for index in rows:
        if sum(1 for value in raw.iloc[index] if _text(value)) >= 2:
            return index
    return 0


def _header_names(raw: pd.DataFrame, header_index: int, merged: dict[int, str]) -> list[str]:
    names = []
    for position, value in enumerate(raw.iloc[header_index]):
        names.append(_text(value) or merged.get(position, ""))
    return names


def _merged_header_names(file_path: Path, sheet_name: str | None, header_index: int) -> dict[int, str]:
    """Return names of header cells covered by a merge that starts above them.

    A merge across several columns names only its first column. Formats that
    openpyxl cannot read (.xls) simply get no merged names.
    """

    try:
        from openpyxl import load_workbook

        workbook = load_workbook(file_path, data_only=True)
        sheet = workbook[sheet_name] if sheet_name else workbook.worksheets[0]
    except Exception:
        return {}

    header_row = header_index + 1
    names: dict[int, str] = {}
    try:
        for merged_range in sheet.merged_cells.ranges:
            if not merged_range.min_row < header_row <= merged_range.max_row:
                continue
            value = _text(sheet.cell(merged_range.min_row, merged_range.min_col).value)
            if value:
                names[merged_range.min_col - 1] = value
    finally:
        workbook.close()
    return names


def _unique(names: list[str]) -> list[str]:
    seen: dict[str, int] = {}
    unique = []
    for name in names:
        count = seen.get(name, 0) + 1
        seen[name] = count
        unique.append(name if count == 1 else f"{name} ({count})")
    return unique


def _text(value: Any) -> str:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ""
    return " ".join(str(value).split())
